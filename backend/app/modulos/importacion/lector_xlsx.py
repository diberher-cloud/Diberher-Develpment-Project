"""Lectura PURA de `.xlsx` con la misma salida que `leer_csv` (§4.9).

Defensas: debe ser un ZIP válido y el tamaño descomprimido declarado no puede exceder el límite
(anti zip-bomb); se abre en modo `read_only` y `data_only` (las fórmulas se leen como su último
valor calculado guardado). openpyxl usa defusedxml automáticamente si está instalado.
"""

import io
import math
import zipfile
from collections.abc import Sequence
from datetime import date, datetime, time
from typing import Any

from app.modulos.importacion.lector_csv import ArchivoCsv, ErrorLector, FilaCsv
from app.modulos.importacion.normalizacion import normalizar_cabecera
from app.saneamiento import limpiar

MAX_DESCOMPRIMIDO_DEFECTO = 100 * 1024 * 1024
_MENSAJE = "No se pudo leer el archivo Excel."


def _invalido(mensaje: str = _MENSAJE) -> ErrorLector:
    return ErrorLector("ARCHIVO_INVALIDO", mensaje)


def celda_a_texto(valor: Any) -> str:
    """Convierte una celda a texto (reglas de §4.9)."""
    if valor is None:
        return ""
    if isinstance(valor, str):
        return valor.strip()
    if isinstance(valor, bool):
        return "true" if valor else "false"
    if isinstance(valor, int):
        return str(valor)
    if isinstance(valor, float):
        if math.isfinite(valor) and valor == int(valor):
            return str(int(valor))
        return repr(valor)
    if isinstance(valor, datetime):
        if valor.time() == time(0, 0) and valor.tzinfo is None:
            return valor.date().isoformat()
        return valor.isoformat()
    if isinstance(valor, date):
        return valor.isoformat()
    if isinstance(valor, time):
        return valor.isoformat()
    return limpiar(str(valor)).strip()


def _verificar_zip(datos: bytes, max_descomprimido: int) -> None:
    try:
        with zipfile.ZipFile(io.BytesIO(datos)) as z:
            total = sum(info.file_size for info in z.infolist())
    except (zipfile.BadZipFile, ValueError, OSError, NotImplementedError) as exc:
        raise _invalido() from exc
    if total > max_descomprimido:
        raise _invalido("El archivo Excel descomprimido excede el tamaño permitido.")


def leer_xlsx(
    datos: bytes,
    max_filas: int,
    requeridas: Sequence[str] = (),
    max_descomprimido: int = MAX_DESCOMPRIMIDO_DEFECTO,
) -> ArchivoCsv:
    _verificar_zip(datos, max_descomprimido)
    try:
        from openpyxl import load_workbook

        libro = load_workbook(io.BytesIO(datos), read_only=True, data_only=True)
    except Exception as exc:
        raise _invalido() from exc

    archivo: ArchivoCsv | None = None
    indices: dict[str, int] = {}
    try:
        hoja = libro.worksheets[0]
        hoja.reset_dimensions()
        for numero, fila in enumerate(hoja.iter_rows(values_only=True), start=1):
            valores = [celda_a_texto(c) for c in fila]
            if not any(valores):
                continue  # fila totalmente vacía
            if archivo is None:
                cabeceras = [normalizar_cabecera(c) for c in valores]
                for i, nombre in enumerate(cabeceras):
                    if nombre and nombre not in indices:
                        indices[nombre] = i
                faltantes = [c for c in requeridas if c not in indices]
                if faltantes:
                    raise ErrorLector(
                        "COLUMNAS_FALTANTES",
                        "Faltan columnas obligatorias.",
                        [
                            {"campo": c, "mensaje": "Columna obligatoria ausente."}
                            for c in faltantes
                        ],
                    )
                archivo = ArchivoCsv(cabeceras=list(indices))
                continue
            if len(archivo.filas) >= max_filas:
                raise ErrorLector(
                    "DEMASIADAS_FILAS", f"El archivo supera el máximo de {max_filas} filas."
                )
            archivo.filas.append(
                FilaCsv(
                    numero=numero,
                    valores={
                        nombre: (valores[i] if i < len(valores) else "")
                        for nombre, i in indices.items()
                    },
                )
            )
    except ErrorLector:
        raise
    except Exception as exc:
        raise _invalido() from exc
    finally:
        libro.close()

    if archivo is None or not archivo.filas:
        raise ErrorLector("ARCHIVO_VACIO", "El archivo no contiene filas de datos.")
    return archivo
